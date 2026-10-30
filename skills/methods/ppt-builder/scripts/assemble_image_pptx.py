#!/usr/bin/env python3
"""Assemble generated full-slide images into one viewable PPTX."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.util import Inches


WIDESCREEN_WIDTH = Inches(13.333333)
WIDESCREEN_HEIGHT = Inches(7.5)


def assemble_presentation(
    image_paths: Iterable[Path | str], output_path: Path | str, *, title: str = "Presentation"
) -> Path:
    paths = [Path(path) for path in image_paths]
    if not paths:
        raise ValueError("at least one slide image is required")
    missing = [str(path) for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"slide images not found: {', '.join(missing)}")

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    presentation.slide_width = WIDESCREEN_WIDTH
    presentation.slide_height = WIDESCREEN_HEIGHT
    presentation.core_properties.title = title
    blank_layout = presentation.slide_layouts[6]
    for image_path in paths:
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(image_path), 0, 0, WIDESCREEN_WIDTH, WIDESCREEN_HEIGHT)

    presentation.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("images", nargs="+", help="Ordered full-slide image paths.")
    parser.add_argument("--output", required=True, help="Output PPTX path.")
    parser.add_argument("--title", default="Presentation", help="Presentation title metadata.")
    args = parser.parse_args()

    assemble_presentation(args.images, args.output, title=args.title)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

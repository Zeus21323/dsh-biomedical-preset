import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from formula_numbering import _add_number_to_paragraph, number_equations_in_docx
from latex_omml_converter import (
    M_NS,
    W_NS,
    find_known_omml_schema_issues,
    find_unconverted_latex_commands,
    normalize_latex_input,
)


class FormulaRegressionTests(unittest.TestCase):
    def test_arg_extrema_are_normalized_before_conversion(self):
        self.assertEqual(r"\arg\max_{x}", normalize_latex_input(r"\argmax_{x}"))
        self.assertEqual(r"\arg\min_{x}", normalize_latex_input(r"\argmin_{x}"))

    def test_literal_latex_in_omml_is_reported(self):
        omml = f'<m:oMath xmlns:m="{M_NS}"><m:r><m:t>\\argmax</m:t></m:r></m:oMath>'
        self.assertEqual([r"\argmax"], find_unconverted_latex_commands(omml))
        self.assertIn("literal LaTeX command in OMML: \\argmax", find_known_omml_schema_issues(omml))

    def test_equation_number_run_has_one_properties_element(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        _add_number_to_paragraph(paragraph, "(1)")
        runs = paragraph._element.findall(f"{{{W_NS}}}r")
        number_run = runs[-1]
        self.assertEqual(1, len(number_run.findall(f"{{{W_NS}}}rPr")))
        self.assertEqual(f"{{{W_NS}}}rPr", number_run[0].tag)
        self.assertEqual([], find_known_omml_schema_issues(paragraph._element))

    def test_only_display_equations_are_numbered(self):
        doc = Document()
        inline = doc.add_paragraph("Inline: ")
        inline._element.append(
            parse_xml(f'<m:oMath xmlns:m="{M_NS}"><m:r><m:t>x</m:t></m:r></m:oMath>')
        )
        display = doc.add_paragraph()
        display._element.append(
            parse_xml(
                f'<m:oMathPara xmlns:m="{M_NS}"><m:oMath><m:r><m:t>x=1</m:t></m:r>'
                f'</m:oMath></m:oMathPara>'
            )
        )

        self.assertEqual(1, number_equations_in_docx(doc))
        self.assertNotIn("(1)", inline.text)
        self.assertIn("(1)", display.text)


if __name__ == "__main__":
    unittest.main()
